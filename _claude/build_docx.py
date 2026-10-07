"""The cheat sheet as a Word document: A4, two sides, three columns, in colour.

Reads the same glossary.py as the web page and the question hints, and the same
print_data.py cuts as the printable web version.
"""
import os, re, sys
HERE = os.path.dirname(os.path.abspath(__file__)); REPO = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from glossary import GROUPS
from print_data import DROP_GROUPS, DROP, SHORT, MAP, TWICE

from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GREEN  = RGBColor(0x1A, 0x64, 0x50)
VIOLET = RGBColor(0x5B, 0x3F, 0x86)
AMBER  = RGBColor(0x8A, 0x5D, 0x12)
INK    = RGBColor(0x17, 0x20, 0x1B)
SOFT   = RGBColor(0x3C, 0x4A, 0x42)
MUTE   = RGBColor(0x6D, 0x7A, 0x72)

P1, P2 = 9.0, 7.1          # found by rendering to PDF and measuring how much of the page filled

def shade(p, hexfill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear"); el.set(qn("w:color"), "auto"); el.set(qn("w:fill"), hexfill)
    p._p.get_or_add_pPr().append(el)

def left_bar(p, hexcol):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "4"); left.set(qn("w:color"), hexcol)
    bdr.append(left); pPr.append(bdr)

def columns(section, n, space_twips=200):
    cols = section._sectPr.xpath("./w:cols")[0]
    cols.set(qn("w:num"), str(n)); cols.set(qn("w:space"), str(space_twips))
    cols.set(qn("w:equalWidth"), "1")

TOKEN = re.compile(r"(<b>|</b>|<i>|</i>)")
def runs(p, html, size, colour=SOFT, base_bold=False, base_ital=False):
    html = (html.replace("&middot;", "·").replace("&ldquo;", "“")
                .replace("&rdquo;", "”").replace("&nbsp;", " ").replace("&amp;", "&"))
    html = re.sub(r"<br\s*/?>", "  ", html)
    bold, ital = base_bold, base_ital
    for part in TOKEN.split(html):
        if part == "<b>":  bold = True;  continue
        if part == "</b>": bold = base_bold;  continue
        if part == "<i>":  ital = True;  continue
        if part == "</i>": ital = base_ital; continue
        part = re.sub(r"<[^>]+>", "", part)
        if not part: continue
        r = p.add_run(part)
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = ital
        r.font.color.rgb = colour; r.font.name = "Calibri"

def para(doc, space_after=0, space_before=0, keep=True, with_next=False):
    """with_next holds a card together, so a heading never ends a column alone."""
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_after = Pt(space_after); f.space_before = Pt(space_before)
    f.line_spacing = 1.0; f.keep_together = keep; f.keep_with_next = with_next
    return p

def build():
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(P1)
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(s, m, Mm(7))
    columns(s, 3)

    head = para(doc, space_after=4)
    runs(head, "PHI2394 · exam 1 · the map", P1 + 2.4, GREEN, base_bold=True)

    for n, name, yr, q, lines, lead, critic in MAP:
        col = VIOLET if critic else GREEN
        hexc = "5B3F86" if critic else "1A6450"
        fill = "EFEAF7" if critic else "E8F1EC"
        h = para(doc, space_before=1, with_next=True); shade(h, fill); left_bar(h, hexc)
        runs(h, "%d  " % n, P1 - 0.8, col, base_bold=True)
        runs(h, name, P1 + 1.1, col, base_bold=True)
        runs(h, "   " + yr, P1 - 1.2, MUTE)
        qp = para(doc, with_next=True); shade(qp, fill); left_bar(qp, hexc)
        runs(qp, q, P1 - 0.3, SOFT, base_ital=True)
        for k, ln in enumerate(lines):
            b = para(doc, with_next=(k < len(lines) - 1)); shade(b, fill); left_bar(b, hexc)
            b.paragraph_format.left_indent = Pt(6)
            runs(b, "• " + ln, P1, SOFT)
        if lead:
            lp = para(doc, space_after=4, space_before=1)
            runs(lp, "↓  " + lead, P1 - 0.8, MUTE, base_ital=True)
        else:
            para(doc, space_after=4)

    th = para(doc, space_before=3, space_after=2, with_next=True); shade(th, "1A6450")
    runs(th, " Who turns up twice", P1 + 1.1, RGBColor(0xFF, 0xFF, 0xFF), base_bold=True)
    for name, times, body in TWICE:
        tp = para(doc, space_after=3); shade(tp, "F6F3EC"); left_bar(tp, "8A5D12")
        runs(tp, name + " ", P1 + 0.4, AMBER, base_bold=True)
        runs(tp, times + "  ", P1 - 1.2, MUTE)
        runs(tp, body, P1, SOFT)

    doc.add_page_break()
    g = para(doc, space_after=4)
    runs(g, "PHI2394 · exam 1 · the glossary", P2 + 3.0, GREEN, base_bold=True)

    for gid, title, rows in GROUPS:
        if gid in DROP_GROUPS: continue
        rows = [r for r in rows if r[0] not in DROP]
        if not rows: continue
        hp = para(doc, space_before=3, space_after=2, with_next=True); shade(hp, "1A6450")
        runs(hp, " " + title, P2 + 1.4, RGBColor(0xFF, 0xFF, 0xFF), base_bold=True)
        for term, who, short, extra in rows:
            s_, e_ = SHORT.get(term, (short, extra))
            tp = para(doc, space_after=2)
            runs(tp, re.sub(r"<[^>]+>", "", term) + " ", P2 + 0.3, GREEN, base_bold=True)
            if who: runs(tp, who + " ", P2 - 0.8, MUTE)
            body = s_[0].upper() + s_[1:] + "."
            if e_: body += " " + e_
            runs(tp, body, P2, SOFT)

    out = os.path.join(REPO, "exam-1", "PHI2394 exam 1 cheat sheet.docx")
    doc.save(out)
    return out

if __name__ == "__main__":
    p = build()
    print("saved", os.path.basename(p), os.path.getsize(p), "bytes")
